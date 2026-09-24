#!/usr/bin/env python3
import io, json, os, re
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

import requests
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

# Reuse the official EIA + FX helpers already tested for Brent.
from brent_weekly import weekly_brent, weekly_fx

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
DOCS.mkdir(exist_ok=True)
OUT_IMG = DOCS / "story-weekly-fuels.jpg"
OUT_JSON = DOCS / "story-weekly-fuels.json"

EU_HISTORY_XLSX = (
    "https://energy.ec.europa.eu/document/download/"
    "906e60ca-8b6a-44e7-8589-652854d2fd3f_en"
    "?filename=Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx"
)
BARREL_LITRES = 158.987294928

BG = (247, 244, 237)
INK = (15, 15, 15)
MUTED = (83, 82, 79)
GREEN = (18, 133, 55)
BLUE = (28, 119, 211)
BLACK = (42, 42, 42)
OLIVE = (99, 111, 79)
CARD = (253, 251, 247)
LIGHT = (245, 243, 237)

def font(size, bold=False, italic=False):
    base = "/usr/share/fonts/truetype/dejavu/DejaVuSans"
    if bold and italic:
        p = base + "-BoldOblique.ttf"
    elif bold:
        p = base + "-Bold.ttf"
    elif italic:
        p = base + "-Oblique.ttf"
    else:
        p = base + ".ttf"
    return ImageFont.truetype(p, size)

def fnum(x, n=3):
    return f"{x:.{n}f}".replace(".", ",")

def to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):
        # Excel serial date
        return (datetime(1899, 12, 30) + timedelta(days=float(v))).date()
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except Exception:
            pass
    return None

def fetch_eu_history():
    r = requests.get(
        EU_HISTORY_XLSX,
        timeout=90,
        headers={"User-Agent": "BermeEnergyAutomation/1.0"},
    )
    r.raise_for_status()
    wb = load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
    if "Prices with taxes" not in wb.sheetnames:
        raise RuntimeError(f"No existe la hoja 'Prices with taxes'. Hojas: {wb.sheetnames}")
    ws = wb["Prices with taxes"]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 5:
        raise RuntimeError("El histórico de carburantes está vacío")

    header = [str(x or "").strip() for x in rows[0]]
    ctr_cols = [i for i, x in enumerate(header) if x == "CTR"]
    if not ctr_cols:
        raise RuntimeError("No se encontraron bloques CTR en el histórico")

    for block_idx, start in enumerate(ctr_cols):
        end = (ctr_cols[block_idx + 1] - 1) if block_idx + 1 < len(ctr_cols) else len(header) - 1

        country_code = None
        for row in rows[3:]:
            v = str(row[start] or "").strip().replace("_", "")
            if v:
                country_code = v
                break
        if country_code != "ES":
            continue

        gas_idx = None
        diesel_idx = None
        for j in range(start + 1, end + 1):
            h = header[j].lower()
            if gas_idx is None and re.search(r"euro.?super\s*95", h):
                gas_idx = j
            if diesel_idx is None and re.search(r"gas\s*oil|gasoil|automotive", h):
                diesel_idx = j

        if gas_idx is None or diesel_idx is None:
            raise RuntimeError("No se encontraron las columnas Gasolina 95 y Diésel A para España")

        data = []
        for row in rows[3:]:
            d = to_date(row[0])
            if not d:
                continue
            try:
                g_raw = float(row[gas_idx]) if row[gas_idx] not in (None, "") else None
                d_raw = float(row[diesel_idx]) if row[diesel_idx] not in (None, "") else None
            except Exception:
                continue
            if g_raw and d_raw:
                # Source is EUR / 1000 litres.
                data.append({"date": d, "gasoline": g_raw / 1000.0, "diesel": d_raw / 1000.0})

        if not data:
            raise RuntimeError("No se encontraron precios históricos de España")
        # Remove duplicate dates and sort.
        dedup = {x["date"]: x for x in data}
        return [dedup[k] for k in sorted(dedup)]

    raise RuntimeError("No se encontró el bloque de España (ES) en el histórico")

def completed_week():
    ref = os.getenv("BERME_REFERENCE_DATE")
    today = date.fromisoformat(ref) if ref else datetime.now(ZoneInfo("Europe/Madrid")).date()
    mon = today - timedelta(days=today.weekday())
    return mon - timedelta(days=7), mon - timedelta(days=1)

def brent_two_weeks():
    cur_start, cur_end = completed_week()
    start = cur_start - timedelta(days=7)
    b = weekly_brent(start, cur_end)
    fx = weekly_fx(start, cur_end)
    common = sorted(set(b) & set(fx))
    values = [(d, (b[d] / fx[d]) / BARREL_LITRES) for d in common]

    cur = [(d, v) for d, v in values if cur_start <= d <= cur_end]
    prev_start = cur_start - timedelta(days=7)
    prev_end = cur_start - timedelta(days=1)
    prev = [(d, v) for d, v in values if prev_start <= d <= prev_end]
    if not cur:
        raise RuntimeError("No hay datos Brent para la última semana completa")
    return cur, prev

def historical_brent_eur_l():
    # EIA Europe Brent spot maximum in 2022 was USD 133.18/bbl on 8 Mar 2022.
    # Convert that reference day to EUR using the same FX service as the weekly series.
    d = date(2022, 3, 8)
    try:
        fx = weekly_fx(d, d)
        usd_per_eur = fx[d]
        return (133.18 / usd_per_eur) / BARREL_LITRES, d
    except Exception:
        return 0.75, d

def build():
    history = fetch_eu_history()
    latest = history[-1]
    previous = history[-2] if len(history) >= 2 else latest

    # Keep a compact trend for the approved template.
    recent = history[-8:]

    g_hist = max(history, key=lambda x: x["gasoline"])
    d_hist = max(history, key=lambda x: x["diesel"])

    cur_brent, prev_brent = brent_two_weeks()
    brent_avg = sum(v for _, v in cur_brent) / len(cur_brent)
    brent_prev = (sum(v for _, v in prev_brent) / len(prev_brent)) if prev_brent else brent_avg
    brent_hist_val, brent_hist_date = historical_brent_eur_l()
    if brent_avg > brent_hist_val:
        brent_hist_val = brent_avg
        brent_hist_date = max(cur_brent, key=lambda x: x[1])[0]

    observed = latest["date"]
    week_start = observed - timedelta(days=6)

    return {
        "week_start": week_start.isoformat(),
        "week_end": observed.isoformat(),
        "gasoline_eur_l": latest["gasoline"],
        "diesel_eur_l": latest["diesel"],
        "brent_eur_l_equiv": brent_avg,
        "gasoline_change_pct": (latest["gasoline"] / previous["gasoline"] - 1) * 100,
        "diesel_change_pct": (latest["diesel"] / previous["diesel"] - 1) * 100,
        "brent_change_pct": (brent_avg / brent_prev - 1) * 100 if brent_prev else 0,
        "gasoline_historical_max": g_hist["gasoline"],
        "gasoline_historical_date": g_hist["date"].isoformat(),
        "diesel_historical_max": d_hist["diesel"],
        "diesel_historical_date": d_hist["date"].isoformat(),
        "brent_historical_max": brent_hist_val,
        "brent_historical_date": brent_hist_date.isoformat(),
        "gasoline_trend": [{"date": x["date"].isoformat(), "value": x["gasoline"]} for x in recent],
        "diesel_trend": [{"date": x["date"].isoformat(), "value": x["diesel"]} for x in recent],
        "brent_trend": [{"date": d.isoformat(), "value": v} for d, v in cur_brent],
        "source": "Comisión Europea Weekly Oil Bulletin (España) + EIA Brent + FX EUR/USD",
    }

def draw_nozzle(img):
    # Fixed vector illustration; never changes week to week.
    d = ImageDraw.Draw(img, "RGBA")
    # hose/body
    d.rounded_rectangle((825, 250, 1095, 480), radius=58, fill=(8, 112, 50, 255))
    d.polygon([(840, 430), (975, 430), (930, 625), (850, 605)], fill=(18, 28, 25, 255))
    # metallic spout
    d.rounded_rectangle((680, 225, 910, 292), radius=28, fill=(150, 154, 150, 255))
    d.rounded_rectangle((650, 237, 730, 280), radius=20, fill=(55, 57, 55, 255))
    d.line((700, 238, 890, 238), fill=(225, 225, 218, 210), width=5)
    # connector rings
    for x in (875, 895, 915):
        d.ellipse((x - 22, 220, x + 22, 300), outline=(72, 74, 72, 255), width=6)

def sparkline(draw, box, series, color):
    x0, y0, x1, y1 = box
    vals = [p["value"] for p in series]
    if not vals:
        return
    lo, hi = min(vals), max(vals)
    pad = max((hi - lo) * 0.2, max(abs(hi), 1) * 0.01)
    lo -= pad
    hi += pad
    if hi <= lo:
        hi = lo + 1
    pts = []
    n = len(vals)
    for i, v in enumerate(vals):
        x = x0 + (x1 - x0) * (i / max(1, n - 1))
        y = y1 - (v - lo) / (hi - lo) * (y1 - y0)
        pts.append((x, y))
    draw.line(pts, fill=color, width=5, joint="curve")
    for x, y in pts:
        draw.ellipse((x - 3, y - 3, x + 3, y + 3), fill=color)

def date_label(iso):
    d = date.fromisoformat(iso)
    months = ["ene","feb","mar","abr","may","jun","jul","ago","sep","oct","nov","dic"]
    return f"{d.day} {months[d.month-1]} {d.year}"

def render(data):
    W, H = 1080, 1920
    img = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(img)

    # Fixed background/layout.
    d.ellipse((775, -130, 1200, 300), fill=(238, 240, 232))
    draw_nozzle(img)

    # Instagram top UI safe area: no important content above y=205.
    s = date.fromisoformat(data["week_start"])
    e = date.fromisoformat(data["week_end"])
    months = ["ENE","FEB","MAR","ABR","MAY","JUN","JUL","AGO","SEP","OCT","NOV","DIC"]
    d.multiline_text(
        (920, 215),
        f"SEMANA DEL\n{s.day} AL {e.day} {months[e.month-1]} {e.year}",
        font=font(23), fill=INK, anchor="ra", align="right", spacing=6
    )

    d.text((72, 330), "Precios medios", font=font(62, bold=True), fill=INK)
    d.text((72, 405), "de", font=font(62, bold=True), fill=INK)
    d.text((180, 405), "combustibles", font=font(62, bold=True), fill=OLIVE)
    d.text((75, 485), "€/litro (España)", font=font(31), fill=INK)

    # Main white card, matching the approved template proportions.
    d.rounded_rectangle((60, 575, 1020, 1455), radius=34, fill=CARD)

    rows = [
        ("Gasolina 95", data["gasoline_eur_l"], data["gasoline_change_pct"], data["gasoline_trend"], GREEN),
        ("Diésel A", data["diesel_eur_l"], data["diesel_change_pct"], data["diesel_trend"], BLUE),
        ("Petróleo Brent", data["brent_eur_l_equiv"], data["brent_change_pct"], data["brent_trend"], BLACK),
    ]
    row_y = [615, 835, 1055]

    for idx, (label, value, pct, trend, color) in enumerate(rows):
        y = row_y[idx]
        d.text((105, y), label, font=font(31, bold=True), fill=INK)
        d.text((105, y + 55), fnum(value, 3), font=font(73, bold=True), fill=OLIVE)
        d.text((365, y + 98), "€/L", font=font(31), fill=INK)

        badge_fill = (235, 241, 230) if pct <= 0 else (246, 236, 225)
        d.rounded_rectangle((445, y + 28, 655, y + 145), radius=22, fill=badge_fill)
        arrow = "↓" if pct <= 0 else "↑"
        pct_color = OLIVE if pct <= 0 else (145, 84, 35)
        d.text((472, y + 43), f"{arrow} {pct:+.0f}%", font=font(37, bold=True), fill=pct_color)
        d.text((472, y + 98), "vs. semana anterior", font=font(17), fill=MUTED)

        d.rounded_rectangle((680, y + 24, 960, y + 145), radius=22, fill=LIGHT)
        sparkline(d, (705, y + 48, 930, y + 110), trend, color)

        if idx < 2:
            # National weekly bulletin has one weekly observation; show current week range as the last
            # two bulletin points rather than inventing intra-week min/max.
            vals = [x["value"] for x in trend[-2:]] or [value]
        else:
            vals = [x["value"] for x in trend] or [value]
        d.text(
            (700, y + 155),
            f"Mín. {fnum(min(vals),3)}  |  Máx. {fnum(max(vals),3)}",
            font=font(18), fill=MUTED
        )

        if idx < 2:
            d.line((105, y + 200, 965, y + 200), fill=(225, 221, 214), width=2)

    # Historical maxima.
    d.rounded_rectangle((95, 1285, 985, 1415), radius=24, fill=LIGHT)
    d.text((125, 1308), "Máximos históricos", font=font(25, bold=True), fill=INK)
    hist_rows = [
        ("Gasolina 95", data["gasoline_historical_max"], data["gasoline_historical_date"], GREEN),
        ("Diésel A", data["diesel_historical_max"], data["diesel_historical_date"], BLUE),
        ("Brent equiv.", data["brent_historical_max"], data["brent_historical_date"], BLACK),
    ]
    hy = [1350, 1382, 1414]
    # Slightly overlap lower edge intentionally? No: compress third row above 1410.
    hy = [1344, 1374, 1404]
    for (label, value, dt, color), y in zip(hist_rows, hy):
        d.ellipse((125, y - 8, 145, y + 12), fill=color)
        d.text((160, y - 10), label, font=font(19), fill=INK)
        d.text((545, y - 10), f"{fnum(value,3)} €/L", font=font(19, bold=True), fill=INK, anchor="ma")
        d.text((935, y - 10), date_label(dt), font=font(18), fill=MUTED, anchor="ra")

    # Source / clarification.
    d.ellipse((83, 1483, 119, 1519), outline=INK, width=3)
    d.text((101, 1482), "i", font=font(23, bold=True), fill=INK, anchor="ma")
    d.text((140, 1474), "Fuente: Comisión Europea (Gasolina 95 y Diésel A) + EIA (Brent).", font=font(18), fill=MUTED)
    d.text((140, 1504), "El Brent es crudo: se muestra su equivalente €/L, no un precio de gasolinera.", font=font(18), fill=MUTED)

    # CTA and signature, both above the Instagram reply bar.
    d.rounded_rectangle((75, 1565, 1005, 1675), radius=26, fill=(94, 109, 70))
    d.text((540, 1593), "Si quieres revisar tu factura, escríbenos.", font=font(26, bold=True), fill="white", anchor="ma")
    d.text((540, 1633), "Te ayudamos gratuitamente.", font=font(23), fill="white", anchor="ma")

    d.line((80, 1725, 155, 1725), fill=GREEN, width=4)
    d.text((80, 1743), "Ahorra con", font=font(29), fill=INK)
    d.text((80, 1780), "Berme Energy.", font=font(31, italic=True), fill=GREEN)

    # Bottom ~95 px plus Instagram UI overlay space remains visually empty.
    img.save(OUT_IMG, "JPEG", quality=94, optimize=True, progressive=True)

def main():
    data = build()
    OUT_JSON.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    render(data)
    print(json.dumps({
        "week": [data["week_start"], data["week_end"]],
        "gasoline_eur_l": round(data["gasoline_eur_l"], 3),
        "diesel_eur_l": round(data["diesel_eur_l"], 3),
        "brent_eur_l_equiv": round(data["brent_eur_l_equiv"], 3),
        "image": str(OUT_IMG),
    }, ensure_ascii=False))

if __name__ == "__main__":
    main()
